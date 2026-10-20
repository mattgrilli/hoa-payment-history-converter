import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import os
import re
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Output columns, in order, with their Excel widths
COLUMNS = [
    ('UNIT', 15), ('ACCT', 10), ('LOT', 10), ('OWNER', 25), ('OWNER TYPE', 12),
    ('DELINQUENCY STATUS', 20), ('TXN DATE', 12), ('PAYMENTS', 12), ('DESCRIPTION', 20),
    ('CHARGES', 12), ('ASSESSMENTS A1', 15), ('LATE FEES 01', 15), ('SPEC ASSMTS 08', 15),
    ('PREPAIDS PP', 15), ('CODE', 10), ('AMOUNT', 12), ('BALANCE', 12), ('NOTES', 40),
]
HEADERS = [name for name, _ in COLUMNS]
OUTPUT_EXTENSIONS = ('.pdf', '.xlsx', '.csv')

def safe_sheet_title(name):
    """Excel sheet titles can't contain : \\ / ? * [ ] and max out at 31 chars"""
    return (re.sub(r'[:\\/?*\[\]]', '-', name).strip("'") or 'Sheet')[:31]

def write_excel_headers(ws):
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

def finish_excel_sheet(ws):
    for col, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

class ReportConverter:
    def __init__(self, root):
        self.root = root
        self.root.title("HOA Payment History Converter")
        self.root.geometry("700x750")
        
        self.style = ttk.Style(self.root)
        try:
            self.style.theme_use('clam') 
        except tk.TclError:
            print("Clam theme not available, using default.")

        self.style.configure('Accent.TButton', 
                             foreground='white', 
                             background='#2196F3', 
                             font=('Arial', 12))
        self.style.map('Accent.TButton',
            background=[('active', '#1976D2')])

        self.style.configure('Success.TButton', 
                             foreground='white', 
                             background='#4CAF50', 
                             font=('Arial', 12, 'bold'))
        self.style.map('Success.TButton',
            background=[('active', '#45a049')])

        # MODIFIED: New style for the bold checkbutton
        self.style.configure('Bold.TCheckbutton', font=('Arial', 11, 'bold'))

        # Variables
        self.input_file = tk.StringVar()
        self.output_file = tk.StringVar()
        self.output_format = tk.StringVar(value="pdf")
        self.font_size = tk.IntVar(value=8)
        self.preserve_formatting = tk.BooleanVar(value=True)
        self.split_by_address = tk.BooleanVar(value=False)
        
        self.single_account_mode = tk.BooleanVar(value=False)
        self.selected_account = tk.StringVar()
        self.available_accounts = []
        self.owner_type_filter = tk.StringVar(value="both")  # "current", "previous", or "both"

        self.create_widgets()
        
    def create_widgets(self):
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.pack(fill="both", expand=True)

        title = ttk.Label(main_frame, text="HOA Payment History Converter", 
                         font=("Arial", 18, "bold"))
        title.pack(pady=(0, 20))
        
        input_frame = ttk.Frame(main_frame)
        input_frame.pack(pady=10, fill="x")
        
        ttk.Label(input_frame, text="Input Text File:", 
                 font=("Arial", 12)).pack(anchor="w")
        
        input_entry_frame = ttk.Frame(input_frame)
        input_entry_frame.pack(fill="x", pady=5)
        
        ttk.Entry(input_entry_frame, textvariable=self.input_file, 
                 font=("Arial", 10), state="readonly").pack(side="left", 
                                                            fill="x", expand=True)
        ttk.Button(input_entry_frame, text="Browse...", 
                  command=self.browse_input).pack(side="right", padx=(5,0))
        
        format_frame = ttk.Frame(main_frame)
        format_frame.pack(pady=10, fill="x")
        
        ttk.Label(format_frame, text="Output Format:", 
                 font=("Arial", 12)).pack(anchor="w")
        
        format_buttons = ttk.Frame(format_frame)
        format_buttons.pack(fill="x", pady=5)
        
        ttk.Radiobutton(format_buttons, text="PDF", variable=self.output_format, 
                       value="pdf", command=self.update_output_extension).pack(side="left", padx=10)
        ttk.Radiobutton(format_buttons, text="Excel (.xlsx)", variable=self.output_format, 
                       value="xlsx", command=self.update_output_extension).pack(side="left", padx=10)
        ttk.Radiobutton(format_buttons, text="CSV", variable=self.output_format, 
                       value="csv", command=self.update_output_extension).pack(side="left", padx=10)
        
        single_account_frame = ttk.Frame(main_frame)
        single_account_frame.pack(pady=10, padx=0, fill="x")
        
        # MODIFIED: Removed 'font' and added 'style'
        self.single_account_check = ttk.Checkbutton(single_account_frame, 
                      text="Process only a single account",
                      variable=self.single_account_mode,
                      style='Bold.TCheckbutton',
                      command=self.toggle_single_account_mode)
        self.single_account_check.pack(anchor="w")

        self.account_selector_frame = ttk.Frame(single_account_frame)
        self.account_selector_frame.pack(fill="x", padx=20, pady=5)

        ttk.Label(self.account_selector_frame, text="Select Account:").pack(side="left", padx=(0, 10))
        self.account_combobox = ttk.Combobox(self.account_selector_frame,
                                             textvariable=self.selected_account,
                                             state="disabled")
        self.account_combobox.pack(side="left", fill="x", expand=True)

        # Owner type filter
        self.owner_filter_frame = ttk.Frame(single_account_frame)
        self.owner_filter_frame.pack(fill="x", padx=20, pady=5)

        ttk.Label(self.owner_filter_frame, text="Show transactions for:").pack(side="left", padx=(0, 10))
        self.owner_current_radio = ttk.Radiobutton(self.owner_filter_frame, text="Current Owner",
                                                    variable=self.owner_type_filter,
                                                    value="current", state="disabled",
                                                    command=self.update_naming_pattern)
        self.owner_current_radio.pack(side="left", padx=5)
        self.owner_previous_radio = ttk.Radiobutton(self.owner_filter_frame, text="Previous Owner",
                                                     variable=self.owner_type_filter,
                                                     value="previous", state="disabled",
                                                     command=self.update_naming_pattern)
        self.owner_previous_radio.pack(side="left", padx=5)
        self.owner_both_radio = ttk.Radiobutton(self.owner_filter_frame, text="Both",
                                                variable=self.owner_type_filter,
                                                value="both", state="disabled",
                                                command=self.update_naming_pattern)
        self.owner_both_radio.pack(side="left", padx=5)

        split_frame = ttk.Frame(main_frame)
        split_frame.pack(pady=10, padx=0, fill="x")
        
        self.split_check = ttk.Checkbutton(split_frame, 
                      text="Create separate files for each address (organized in folders)",
                      variable=self.split_by_address,
                      style='Bold.TCheckbutton', # Also applied style here for consistency
                      command=self.toggle_split_mode)
        self.split_check.pack(anchor="w")
        
        ttk.Label(split_frame, 
                 text="When enabled, creates a main folder with subfolders for each unit",
                 font=("Arial", 9)).pack(anchor="w", padx=20)
        
        output_frame = ttk.Frame(main_frame)
        output_frame.pack(pady=10, fill="x")
        
        self.output_label = ttk.Label(output_frame, text="Output File Name:", 
                 font=("Arial", 12))
        self.output_label.pack(anchor="w")
        
        output_entry_frame = ttk.Frame(output_frame)
        output_entry_frame.pack(fill="x", pady=5)
        
        self.output_entry = ttk.Entry(output_entry_frame, textvariable=self.output_file, 
                 font=("Arial", 10))
        self.output_entry.pack(side="left", fill="x", expand=True)

        self.browse_output_btn = ttk.Button(output_entry_frame, text="Choose Location...",
                  command=self.browse_output)
        self.browse_output_btn.pack(side="right", padx=(5,0))

        # Naming pattern label
        self.naming_pattern_label = ttk.Label(output_frame,
                 text="",
                 font=("Arial", 9),
                 foreground="gray")
        self.naming_pattern_label.pack(anchor="w", padx=5, pady=(0, 5))

        self.options_frame = ttk.LabelFrame(main_frame, text="PDF Options", 
                                      padding=(20, 10))
        self.options_frame.pack(pady=10, fill="x")
        
        font_frame = ttk.Frame(self.options_frame)
        font_frame.pack(fill="x", pady=5)
        ttk.Label(font_frame, text="Font Size:", 
                 font=("Arial", 10)).pack(side="left")
        ttk.Spinbox(font_frame, from_=6, to=12, textvariable=self.font_size,
                   width=5, font=("Arial", 10)).pack(side="left", padx=10)
        
        ttk.Checkbutton(self.options_frame, 
                       text="Preserve fixed-width formatting",
                       variable=self.preserve_formatting,
                       ).pack(anchor="w", pady=5)
        
        info_text = "Excel/CSV formats will parse the report into structured columns."
        self.info_label = ttk.Label(main_frame, text=info_text, 
                 font=("Arial", 9), wraplength=600)
        self.info_label.pack(pady=5)
        
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(pady=10)
        
        preview_btn = ttk.Button(button_frame, text="Preview Parse", 
                                command=self.preview_parse,
                                style='Accent.TButton',
                                padding=(15, 8))
        preview_btn.pack(side="left", padx=5)
        
        convert_btn = ttk.Button(button_frame, text="Convert", 
                                command=self.convert,
                                style='Success.TButton',
                                padding=(20, 8))
        convert_btn.pack(side="left", padx=5)
    
    def toggle_single_account_mode(self):
        if self.single_account_mode.get():
            self.account_combobox.config(state="readonly")
            self.owner_current_radio.config(state="normal")
            self.owner_previous_radio.config(state="normal")
            self.owner_both_radio.config(state="normal")
            self.split_by_address.set(False)
            self.split_check.config(state="disabled")
            self.toggle_split_mode() # Update output label
        else:
            self.account_combobox.config(state="disabled")
            self.owner_current_radio.config(state="disabled")
            self.owner_previous_radio.config(state="disabled")
            self.owner_both_radio.config(state="disabled")
            self.split_check.config(state="enabled")
            self.selected_account.set("")
            self.toggle_split_mode()

    def toggle_split_mode(self):
        is_split_mode = self.split_by_address.get()
        is_single_mode = self.single_account_mode.get()
        self.sync_output_name()

        if is_split_mode and not is_single_mode:
            self.output_label.config(text="Output Folder Name:")
            self.browse_output_btn.config(text="Choose Folder...")
            self.output_entry.config(state="normal")
            self.single_account_mode.set(False)
            self.single_account_check.config(state="disabled")
            self.naming_pattern_label.config(text="Folder structure: [Unit]/Unit_Owner_Current.ext, Unit_Owner_Previous.ext")
        elif is_single_mode:
            self.output_label.config(text="Output File Name:")
            self.browse_output_btn.config(text="Choose Location...")
            self.output_entry.config(state="normal")
            owner_filter = self.owner_type_filter.get()
            if owner_filter == "current":
                self.naming_pattern_label.config(text="Format: Unit_Owner_Current.ext")
            elif owner_filter == "previous":
                self.naming_pattern_label.config(text="Format: Unit_Owner_Previous.ext")
            else:
                self.naming_pattern_label.config(text="Creates 2 files: Unit_Owner_Current.ext, Unit_Owner_Previous.ext")
        else:
            self.output_label.config(text="Output File Name:")
            self.browse_output_btn.config(text="Choose Location...")
            self.output_entry.config(state="normal")
            self.single_account_check.config(state="normal")
            self.naming_pattern_label.config(text="")

    def sync_output_name(self):
        """Switch the output name between a folder name and a file name when split mode changes"""
        current = self.output_file.get()
        if not current:
            return
        base, ext = os.path.splitext(current)
        if self.split_by_address.get() and not self.single_account_mode.get():
            if ext.lower() in OUTPUT_EXTENSIONS:
                self.output_file.set(base + "_Reports")
        elif not ext or ext.lower() in OUTPUT_EXTENSIONS:
            if not ext and base.endswith("_Reports"):
                base = base[:-len("_Reports")]
            self.output_file.set(base + "." + self.output_format.get())

    def update_naming_pattern(self):
        """Update the naming pattern label based on current mode"""
        if self.single_account_mode.get():
            owner_filter = self.owner_type_filter.get()
            if owner_filter == "current":
                self.naming_pattern_label.config(text="Format: Unit_Owner_Current.ext")
            elif owner_filter == "previous":
                self.naming_pattern_label.config(text="Format: Unit_Owner_Previous.ext")
            else:
                self.naming_pattern_label.config(text="Creates 2 files: Unit_Owner_Current.ext, Unit_Owner_Previous.ext")

    def update_output_extension(self):
        current_file = self.output_file.get()
        if current_file and not self.split_by_address.get():
            base = os.path.splitext(current_file)[0]
            ext = "." + self.output_format.get()
            self.output_file.set(base + ext)
        
        if self.output_format.get() == "pdf":
            self.options_frame.pack(pady=10, fill="x", before=self.info_label)
        else:
            self.options_frame.pack_forget()
        
    def browse_input(self):
        filename = filedialog.askopenfilename(
            title="Select Text File",
            filetypes=[("Text files", "*.txt"), ("All files", "*.*")]
        )
        if filename:
            self.input_file.set(filename)
            base_name = os.path.basename(filename)
            name_without_ext = os.path.splitext(base_name)[0]
            
            if self.split_by_address.get():
                self.output_file.set(name_without_ext + "_Reports")
            else:
                ext = "." + self.output_format.get()
                self.output_file.set(name_without_ext + ext)
            
            self.load_accounts_from_file(filename)

    def load_accounts_from_file(self, filename):
        try:
            transactions = self.parse_hoa_report(filename)
            accounts = sorted(list(set(
                txn['UNIT'] for txn in transactions if txn.get('UNIT')
            )))
            self.available_accounts = accounts
            self.account_combobox['values'] = self.available_accounts
            if self.available_accounts:
                self.selected_account.set(self.available_accounts[0])
                messagebox.showinfo("Accounts Loaded", 
                                    f"Found {len(self.available_accounts)} unique accounts.")
            else:
                messagebox.showwarning("No Accounts", "Could not find any accounts in the selected file.")
        except Exception as e:
            messagebox.showerror("Parse Error", f"Could not parse accounts from file:\n{e}")

    def browse_output(self):
        if self.split_by_address.get():
            folder = filedialog.askdirectory(title="Select Output Folder")
            if folder:
                self.output_file.set(folder)
        else:
            current_name = self.output_file.get() or "output"
            format_type = self.output_format.get()
            ext_map = {"pdf": ".pdf", "xlsx": ".xlsx", "csv": ".csv"}
            filetypes = [(f"{format_type.upper()} files", f"*{ext_map[format_type]}"), ("All files", "*.*")]
            
            filename = filedialog.asksaveasfilename(
                title="Save As",
                initialfile=current_name,
                defaultextension=ext_map[format_type],
                filetypes=filetypes
            )
            if filename:
                self.output_file.set(filename)
    
    def preview_parse(self):
        if not self.input_file.get():
            messagebox.showerror("Error", "Please select an input file first!")
            return
        
        try:
            transactions = self.parse_hoa_report(self.input_file.get())
            preview_window = tk.Toplevel(self.root)
            preview_window.title("Parse Preview")
            preview_window.geometry("900x600")
            
            text_widget = scrolledtext.ScrolledText(preview_window, wrap=tk.NONE, font=("Courier", 9))
            text_widget.pack(fill="both", expand=True, padx=10, pady=10)
            
            text_widget.insert("1.0", f"Total transactions found: {len(transactions)}\n")
            text_widget.insert("end", "="*100 + "\n\n")
            text_widget.insert("end", "Showing all transactions:\n\n")

            for i, txn in enumerate(transactions):
                if txn.get('_SEPARATOR'):
                    text_widget.insert("end", "--- ACCOUNT SEPARATOR ---\n\n")
                else:
                    text_widget.insert("end", f"Transaction {i+1}:\n")
                    for key, value in txn.items():
                        if value and not key.startswith('_'):
                            text_widget.insert("end", f"  {key}: {value}\n")
                    text_widget.insert("end", "\n")
            text_widget.config(state="disabled")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to parse file:\n{str(e)}")
    
    def convert(self):
        if not self.input_file.get() or not self.output_file.get():
            messagebox.showerror("Error", "Please select an input file and specify an output location!")
            return
        
        try:
            if self.single_account_mode.get():
                self.convert_single_account()
            elif self.split_by_address.get():
                self.convert_split_by_address()
            else:
                self.convert_single_file()
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to create file(s):\n{str(e)}")

    def convert_single_account(self):
        selected_unit = self.selected_account.get()
        if not selected_unit:
            messagebox.showerror("Error", "Please select an account from the dropdown.")
            return

        all_transactions = self.parse_hoa_report(self.input_file.get())
        units_data = self._group_transactions_by_unit(all_transactions)

        unit_data = units_data.get(selected_unit)
        if not unit_data:
            messagebox.showerror("Error", f"Could not find data for account: {selected_unit}")
            return

        # Filter transactions based on owner type selection
        owner_filter = self.owner_type_filter.get()
        input_dir = os.path.dirname(self.input_file.get())
        format_type = self.output_format.get()
        safe_unit = self.sanitize_filename(selected_unit)

        if owner_filter == "both":
            # Create TWO separate files - one for current, one for previous
            current_transactions = []
            previous_transactions = []
            current_owner_name = None
            previous_owner_name = None

            for txn in unit_data['transactions']:
                owner_type = txn.get('OWNER TYPE', '').upper()
                owner_name = txn.get('OWNER', '')
                if owner_type == "CURRENT":
                    current_transactions.append(txn)
                    if not current_owner_name and owner_name:
                        current_owner_name = owner_name
                elif owner_type == "PREVIOUS":
                    previous_transactions.append(txn)
                    if not previous_owner_name and owner_name:
                        previous_owner_name = owner_name

            created_files = []

            # Create current owner file if there are transactions
            if current_transactions:
                current_unit_data = {
                    'unit': unit_data['unit'],
                    'acct': unit_data['acct'],
                    'lot': unit_data['lot'],
                    'owner': current_owner_name or unit_data['owner'],
                    'owner_type': 'CURRENT',
                    'delinquency': unit_data['delinquency'],
                    'transactions': current_transactions
                }
                safe_owner = self.sanitize_filename(current_owner_name) if current_owner_name else 'Unknown'
                filename = f"{safe_unit}_{safe_owner}_Current.{format_type}"
                output_path = os.path.join(input_dir, filename)

                try:
                    if format_type == "pdf":
                        final_path = self.write_unit_pdf(current_unit_data, output_path)
                    elif format_type == "xlsx":
                        final_path = self.write_unit_excel(current_unit_data, output_path)
                    else:
                        final_path = self.write_unit_csv(current_unit_data, output_path)
                    created_files.append(final_path)
                except Exception as e:
                    messagebox.showerror("Error", f"Failed to write current owner file:\n{str(e)}")
                    return

            # Create previous owner file if there are transactions
            if previous_transactions:
                previous_unit_data = {
                    'unit': unit_data['unit'],
                    'acct': unit_data['acct'],
                    'lot': unit_data['lot'],
                    'owner': previous_owner_name or 'Previous Owner',
                    'owner_type': 'PREVIOUS',
                    'delinquency': unit_data['delinquency'],
                    'transactions': previous_transactions
                }
                safe_owner = self.sanitize_filename(previous_owner_name) if previous_owner_name else 'Unknown'
                filename = f"{safe_unit}_{safe_owner}_Previous.{format_type}"
                output_path = os.path.join(input_dir, filename)

                try:
                    if format_type == "pdf":
                        final_path = self.write_unit_pdf(previous_unit_data, output_path)
                    elif format_type == "xlsx":
                        final_path = self.write_unit_excel(previous_unit_data, output_path)
                    else:
                        final_path = self.write_unit_csv(previous_unit_data, output_path)
                    created_files.append(final_path)
                except Exception as e:
                    messagebox.showerror("Error", f"Failed to write previous owner file:\n{str(e)}")
                    return

            if not created_files:
                messagebox.showwarning("No Data", f"No transactions found for {selected_unit}.")
                return

            files_list = "\n".join(created_files)
            messagebox.showinfo("Success", f"Created {len(created_files)} files for {selected_unit}:\n\n{files_list}")
            return

        # Single file mode (current or previous only)
        filtered_transactions = []
        for txn in unit_data['transactions']:
            owner_type = txn.get('OWNER TYPE', '').upper()
            if owner_filter == "current" and owner_type == "CURRENT":
                filtered_transactions.append(txn)
            elif owner_filter == "previous" and owner_type == "PREVIOUS":
                filtered_transactions.append(txn)
        unit_data['transactions'] = filtered_transactions

        if not filtered_transactions:
            messagebox.showwarning("No Data",
                f"No transactions found for {selected_unit} with {owner_filter} owner filter.")
            return

        # The unit is keyed by whichever owner appears first, so take the
        # owner from the filtered transactions instead
        filtered_owner = next((txn['OWNER'] for txn in filtered_transactions if txn.get('OWNER')), None)
        if filtered_owner:
            unit_data['owner'] = filtered_owner

        # Generate a better default filename based on unit and owner
        output_path = self.output_file.get()
        if not output_path or output_path.endswith('.txt') or not os.path.dirname(output_path):
            safe_owner = self.sanitize_filename(unit_data['owner']) if unit_data['owner'] else 'Unknown'

            suffix = "Current" if owner_filter == "current" else "Previous"
            filename = f"{safe_unit}_{safe_owner}_{suffix}.{format_type}"

            output_path = os.path.join(input_dir, filename)

        format_type = self.output_format.get()
        try:
            if format_type == "pdf":
                final_path = self.write_unit_pdf(unit_data, output_path)
            elif format_type == "xlsx":
                final_path = self.write_unit_excel(unit_data, output_path)
            else: # csv
                final_path = self.write_unit_csv(unit_data, output_path)

            messagebox.showinfo("Success", f"Report for {selected_unit} created successfully!\n\n{final_path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to write file:\n{str(e)}")

    def convert_single_file(self):
        output_path = self.output_file.get()
        if not os.path.dirname(output_path):
            input_dir = os.path.dirname(self.input_file.get())
            output_path = os.path.join(input_dir, output_path)
        
        format_type = self.output_format.get()
        if format_type == "pdf":
            self.text_to_pdf(self.input_file.get(), output_path, self.font_size.get(), self.preserve_formatting.get())
        elif format_type == "xlsx":
            self.text_to_excel(self.input_file.get(), output_path)
        else:
            self.text_to_csv(self.input_file.get(), output_path)
        
        messagebox.showinfo("Success", f"File created successfully!\n\n{output_path}")
    
    def _group_transactions_by_unit(self, all_transactions, separate_by_owner_type=False):
        units_data = {}
        current_unit_data = None

        for txn in all_transactions:
            if txn.get('_SEPARATOR'):
                continue

            if txn.get('UNIT'):
                unit_name = txn.get('UNIT')
                owner_name = txn.get('OWNER', '')
                owner_type = txn.get('OWNER TYPE', '')

                # Create a unique key based on whether we're separating by owner type
                if separate_by_owner_type:
                    # Key includes owner and owner type for separation
                    key = f"{unit_name}|||{owner_name}|||{owner_type}"
                else:
                    # Key is just unit name (old behavior for single account mode)
                    key = unit_name

                # If this is a new key, create entry
                if key not in units_data:
                    current_unit_data = {
                        'unit': unit_name,
                        'acct': txn.get('ACCT', ''),
                        'lot': txn.get('LOT', ''),
                        'owner': owner_name,
                        'owner_type': owner_type,
                        'delinquency': txn.get('DELINQUENCY STATUS', ''),
                        'transactions': []
                    }
                    units_data[key] = current_unit_data
                else:
                    current_unit_data = units_data[key]

            if current_unit_data:
                current_unit_data['transactions'].append(txn)
        return units_data

    def convert_split_by_address(self):
        all_transactions = self.parse_hoa_report(self.input_file.get())
        units_data = self._group_transactions_by_unit(all_transactions, separate_by_owner_type=True)

        output_base = self.output_file.get()
        if not os.path.isabs(output_base):
            input_dir = os.path.dirname(self.input_file.get())
            output_base = os.path.join(input_dir, output_base)

        os.makedirs(output_base, exist_ok=True)

        format_type = self.output_format.get()
        ext = "." + format_type

        created_files = []
        # Group by unit address first to create one folder per address
        address_groups = {}
        for key, unit_data in units_data.items():
            unit_name = unit_data['unit']
            if unit_name not in address_groups:
                address_groups[unit_name] = []
            address_groups[unit_name].append(unit_data)

        # Create folders per address, with separate files for each owner
        for unit_name, unit_list in address_groups.items():
            safe_unit = self.sanitize_filename(unit_name)
            unit_folder = os.path.join(output_base, safe_unit)
            os.makedirs(unit_folder, exist_ok=True)

            for unit_data in unit_list:
                safe_owner = self.sanitize_filename(unit_data['owner']) if unit_data['owner'] else 'Unknown'
                owner_type_label = unit_data.get('owner_type', '').capitalize()

                # Create file name: Unit_Owner_OwnerType.ext
                if owner_type_label:
                    file_name = f"{safe_unit}_{safe_owner}_{owner_type_label}{ext}"
                else:
                    file_name = f"{safe_unit}_{safe_owner}{ext}"

                output_file = os.path.join(unit_folder, file_name)

                if format_type == "pdf":
                    self.write_unit_pdf(unit_data, output_file)
                elif format_type == "xlsx":
                    self.write_unit_excel(unit_data, output_file)
                else:
                    self.write_unit_csv(unit_data, output_file)
                created_files.append(output_file)

        messagebox.showinfo("Success",
                          f"Created {len(created_files)} files in:\n\n{output_base}\n\n" +
                          f"Each address has its own folder with separate files for each owner.")
    
    def write_unit_csv(self, unit_data, output_file):
        """Write CSV for a single unit"""
        def write_csv(filepath):
            with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.DictWriter(csvfile, fieldnames=HEADERS)
                writer.writeheader()
                for transaction in unit_data['transactions']:
                    writer.writerow({k: v for k, v in transaction.items() if not k.startswith('_')})

        # Try to write, handling locked files
        try:
            write_csv(output_file)
            return output_file
        except (PermissionError, IOError):
            # File might be open, try with a suffix
            final_path = self.get_available_filename(output_file)
            write_csv(final_path)
            return final_path
    
    def write_unit_excel(self, unit_data, output_file):
        """Write Excel for a single unit"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = safe_sheet_title(unit_data['unit'])

        write_excel_headers(ws)
        for row_idx, transaction in enumerate(unit_data['transactions'], start=2):
            for col_idx, header in enumerate(HEADERS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=transaction.get(header, ''))
        finish_excel_sheet(ws)

        # Try to save, handling locked files
        try:
            wb.save(output_file)
            return output_file
        except (PermissionError, IOError):
            # File might be open, try with a suffix
            final_path = self.get_available_filename(output_file)
            wb.save(final_path)
            return final_path
    
    def write_unit_pdf(self, unit_data, output_file):
        """Write PDF for a single unit"""
        def build_pdf(filepath):
            doc = SimpleDocTemplate(filepath, pagesize=landscape(letter),
                rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
            elements = []
            styles = getSampleStyleSheet()

            title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, spaceAfter=20)
            mono_style = ParagraphStyle('MonoStyle', parent=styles['Normal'], fontName='Courier', fontSize=8, leading=10)

            elements.append(Paragraph(f"Payment History - {unit_data['unit']}", title_style))
            elements.append(Spacer(1, 0.2*inch))

            info_text = f"Account: {unit_data['acct']} | Owner: {unit_data['owner']}"
            if unit_data['delinquency']: info_text += f" | Status: {unit_data['delinquency']}"
            elements.append(Paragraph(info_text, styles['Normal']))
            elements.append(Spacer(1, 0.3*inch))

            for txn in unit_data['transactions']:
                line_parts = []
                if txn.get('TXN DATE'): line_parts.append(txn['TXN DATE'])
                if txn.get('PAYMENTS'): line_parts.append(f"PMT: {txn['PAYMENTS']}")
                if txn.get('DESCRIPTION'): line_parts.append(txn['DESCRIPTION'])
                if txn.get('BALANCE'): line_parts.append(f"BAL: {txn['BALANCE']}")
                if line_parts: elements.append(Paragraph(" | ".join(line_parts), mono_style))
            doc.build(elements)

        # Try to write, handling locked files
        try:
            build_pdf(output_file)
            return output_file
        except (PermissionError, IOError):
            # File might be open, try with a suffix
            final_path = self.get_available_filename(output_file)
            build_pdf(final_path)
            return final_path
    
    def sanitize_filename(self, name):
        """Convert a string to a safe filename"""
        # Replace problematic characters with underscores
        safe_name = re.sub(r'[<>:"/\\|?*]', '_', name)
        # Replace spaces with underscores
        safe_name = safe_name.replace(' ', '_')
        # Remove multiple consecutive underscores
        safe_name = re.sub(r'_+', '_', safe_name)
        # Remove leading/trailing underscores
        safe_name = safe_name.strip('_')
        return safe_name

    def get_available_filename(self, filepath):
        """Get an available filename by appending (1), (2), etc. if file is locked"""
        if not os.path.exists(filepath):
            return filepath

        # Split into base and extension
        base, ext = os.path.splitext(filepath)
        counter = 1

        # Try up to 100 attempts
        while counter < 100:
            new_path = f"{base} ({counter}){ext}"
            if not os.path.exists(new_path):
                return new_path
            counter += 1

        # If we couldn't find an available name, return original
        return filepath

    def is_valid_transaction_line(self, line):
        if len(line) < 40: return False
        first_8 = line[0:8]
        if re.match(r'^\d{2}/\d{2}/\d{2}$', first_8.strip()): return True
        if first_8.strip() == '':
            if any(keyword in line for keyword in ['APPLY', 'EXPENSE', 'INIT']):
                if re.search(r'\d+\.\d{2}', line): return True
        return False
    
    def parse_hoa_report(self, input_file):
        with open(input_file, 'r', encoding='utf-8') as file: content = file.read()
        content = content.replace('~formfeed~', '')
        lines = content.split('\n')
        transactions = []
        current_unit = None; current_account = None; current_lot = None
        current_owner = None; current_owner_type = None; current_delinquency = None; last_transaction = None
        prev_unit = None; prev_account = None; prev_lot = None
        prev_owner = None; prev_owner_type = None; prev_delinquency = None

        for line in lines:
            if line.startswith('UNIT:'):
                match = re.search(r'UNIT:\s+(.+?)\s+ACCT:\s+(\S+)\s+LOT:\s+(\S+)\s+DELINQUENCY STATUS:\s*(.*)$', line)
                if match:
                    current_unit, current_account, current_lot, current_delinquency = [g.strip() for g in match.groups()]
                    # Don't let a unit without an owner line inherit the previous unit's owner
                    current_owner = ''; current_owner_type = ''
                continue

            if line.startswith('CURRENT OWNER:') or line.startswith('PREVIOUS OWNER:'):
                current_owner_type = 'CURRENT' if line.startswith('CURRENT OWNER:') else 'PREVIOUS'
                owner_label = 'CURRENT OWNER:' if line.startswith('CURRENT OWNER:') else 'PREVIOUS OWNER:'
                current_owner = line.replace(owner_label, '').strip()
                continue
            
            if any(x in line for x in ['TIME:', 'PAGE', 'PAYMENT HISTORY REPORT', 'ASSESSMENTS  LATE FEES', 'TXN                   TRANSACTION', 'DATE     PAYMENTS     DESCRIPTION', 'DATE:', '----']) or not line.strip():
                continue
            
            if not self.is_valid_transaction_line(line):
                if last_transaction and line.strip():
                    note = line.strip()
                    last_transaction['NOTES'] = f"{last_transaction['NOTES']} | {note}" if last_transaction['NOTES'] else note
                continue
            
            txn_date = line[0:8].strip(); payments = line[9:21].strip()
            description = line[21:39].strip(); charges = line[39:51].strip()
            assessments_a1 = line[51:63].strip(); late_fees_01 = line[63:75].strip()
            spec_assmts_08 = line[75:87].strip(); prepaids_pp = line[87:99].strip()
            
            remaining = line[99:].strip(); parts = remaining.split()
            code, amount, balance = '', '', ''
            
            if len(parts) == 1:
                balance = parts[0]
            elif len(parts) == 2:
                code, amount = parts if len(parts[0]) <= 3 else ('', parts[0])
                balance = parts[1] if len(parts[0]) > 3 else ''
            elif len(parts) >= 3:
                code, amount, balance = parts[0], parts[1], parts[2]

            unit_changed = (current_unit != prev_unit)
            owner_changed = (current_owner != prev_owner) or (current_owner_type != prev_owner_type)

            if unit_changed and prev_account is not None:
                transactions.append({'_SEPARATOR': True})

            transaction = {'_SEPARATOR': False, 'NOTES': ''}
            if unit_changed or owner_changed:
                transaction.update({'UNIT': current_unit, 'ACCT': current_account, 'LOT': current_lot,
                                    'OWNER': current_owner, 'OWNER TYPE': current_owner_type,
                                    'DELINQUENCY STATUS': current_delinquency})
                prev_unit, prev_account, prev_lot = current_unit, current_account, current_lot
                prev_owner, prev_owner_type, prev_delinquency = current_owner, current_owner_type, current_delinquency
            else:
                # Keep OWNER TYPE populated for all transactions in the group so filtering works
                transaction.update({'UNIT': '', 'ACCT': '', 'LOT': '', 'OWNER': '', 'OWNER TYPE': current_owner_type, 'DELINQUENCY STATUS': ''})

            transaction.update({'TXN DATE': txn_date, 'PAYMENTS': payments, 'DESCRIPTION': description, 'CHARGES': charges,
                                'ASSESSMENTS A1': assessments_a1, 'LATE FEES 01': late_fees_01, 'SPEC ASSMTS 08': spec_assmts_08,
                                'PREPAIDS PP': prepaids_pp, 'CODE': code, 'AMOUNT': amount, 'BALANCE': balance})
            transactions.append(transaction)
            last_transaction = transaction
        return transactions
    
    def text_to_csv(self, input_file, output_file):
        transactions = self.parse_hoa_report(input_file)
        if not transactions: raise ValueError("No transaction data found")
        with open(output_file, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=HEADERS)
            writer.writeheader()
            for transaction in transactions:
                if not transaction.get('_SEPARATOR'):
                    writer.writerow({k: v for k, v in transaction.items() if not k.startswith('_')})
                else:
                    writer.writerow({k: '' for k in HEADERS})
    
    def text_to_excel(self, input_file, output_file):
        transactions = self.parse_hoa_report(input_file)
        if not transactions: raise ValueError("No transaction data found")
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Payment History"

        write_excel_headers(ws)
        separator_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        for row_idx, transaction in enumerate(transactions, start=2):
            if transaction.get('_SEPARATOR'):
                for col_idx in range(1, len(HEADERS) + 1):
                    ws.cell(row=row_idx, column=col_idx, value='').fill = separator_fill
            else:
                for col_idx, header in enumerate(HEADERS, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=transaction.get(header, ''))
        finish_excel_sheet(ws)
        wb.save(output_file)
    
    def text_to_pdf(self, input_file, output_file, font_size, preserve_formatting):
        doc = SimpleDocTemplate(output_file, pagesize=landscape(letter),
            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        elements = []
        styles = getSampleStyleSheet()
        mono_style = ParagraphStyle('MonoStyle', parent=styles['Normal'], fontName='Courier', fontSize=font_size, leading=font_size + 2)
        
        with open(input_file, 'r', encoding='utf-8') as file: content = file.read()
        
        if preserve_formatting:
            pages = content.split('~formfeed~')
            for i, page in enumerate(pages):
                if not page.strip(): continue
                for line in page.split('\n'):
                    if line.strip(): elements.append(Paragraph(line.replace(' ', '&nbsp;'), mono_style))
                    else: elements.append(Spacer(1, font_size))
                if i < len(pages) - 1: elements.append(PageBreak())
        else:
            for para_text in content.replace('~formfeed~', '\n\n').split('\n\n'):
                if para_text.strip():
                    elements.append(Paragraph(para_text.replace('\n', ' ').strip(), mono_style))
                    elements.append(Spacer(1, 0.2*inch))
        doc.build(elements)

def main():
    root = tk.Tk()
    app = ReportConverter(root)
    root.mainloop()

if __name__ == "__main__":
    main()